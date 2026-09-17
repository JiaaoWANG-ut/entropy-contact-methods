#!/usr/bin/env python3
"""Figures and plot-ready tables for the FIG4 entropy analysis."""

from __future__ import annotations

import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from matplotlib import cm
from matplotlib.colors import Normalize
from matplotlib.lines import Line2D

import run_entropy as RE
from run_entropy import SYSTEMS, TEMPERATURES
from traj_entropy import MET_CUT, O_CUT

FIG_W, FIG_H = 7.0, 3.0  # 21:9

# Output location and window caption, set by build().  They are kept here
# rather than read from run_entropy, because that module is a separate object
# when run_entropy.py is executed as a script.
_OUT_DIR = RE.OUT_DIR
_WINDOW = RE.WINDOW_LABEL
METRICS = {
    "s_atom": ("S_atom", "Atomic entropy  $S_\\mathrm{atom}$  (nats)"),
    "s_config": ("S_config", "Configurational entropy  $S_\\mathrm{config}$  (nats)"),
}


def apply_style():
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
            "font.size": 7,
            "axes.labelsize": 8,
            "axes.titlesize": 8,
            "xtick.labelsize": 7,
            "ytick.labelsize": 7,
            "legend.fontsize": 6.5,
            "axes.linewidth": 0.6,
            "xtick.major.width": 0.6,
            "ytick.major.width": 0.6,
            "xtick.major.size": 3,
            "ytick.major.size": 3,
            "lines.linewidth": 1.0,
            "figure.dpi": 150,
            "savefig.dpi": 600,
            "savefig.bbox": "tight",
            "savefig.pad_inches": 0.02,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
        }
    )


def style_axes(ax, hide_right=True):
    ax.spines["top"].set_visible(False)
    if hide_right:
        ax.spines["right"].set_visible(False)
    ax.tick_params(axis="both", direction="in", length=3, width=0.6)


def temp_color(temp_k):
    norm = Normalize(vmin=min(TEMPERATURES), vmax=max(TEMPERATURES))
    return cm.coolwarm(norm(temp_k))


def series(results, system, temp_k):
    payload = results.get((system, temp_k))
    if payload is None:
        return None
    table = payload["table"]
    order = np.argsort(table["time_ps"])
    return table[order]


def save(fig, stem):
    for ext in ("png", "svg"):
        path = os.path.join(_OUT_DIR, f"{stem}.{ext}")
        fig.savefig(path)
        print(f"Saved: {path}")
    plt.close(fig)


def plot_single(results, system, metric):
    label, ylabel = METRICS[metric]
    apply_style()
    fig, ax = plt.subplots(figsize=(FIG_W, FIG_H))
    for temp_k in TEMPERATURES:
        tab = series(results, system, temp_k)
        if tab is None:
            continue
        ax.plot(tab["time_ps"], tab[metric], color=temp_color(temp_k),
                lw=1.1, label=f"{temp_k} K", zorder=3)
    ax.set_xscale("log")
    ax.set_xlabel("Time (ps)", labelpad=3)
    ax.set_ylabel(ylabel, labelpad=3)
    ax.set_title(f"{system}  |  {label} vs time  ({_WINDOW})",
                 loc="left", pad=8)
    ax.legend(frameon=False, loc="best", handlelength=1.4, borderpad=0.3,
              labelspacing=0.25, ncol=3, columnspacing=1.0)
    style_axes(ax)
    save(fig, f"{system}_{label}_time")


def plot_dual(results, system):
    apply_style()
    fig, ax1 = plt.subplots(figsize=(4.8, 2.9))
    ax2 = ax1.twinx()
    for temp_k in TEMPERATURES:
        tab = series(results, system, temp_k)
        if tab is None:
            continue
        color = temp_color(temp_k)
        ax1.plot(tab["time_ps"], tab["s_atom"], color=color, lw=1.1, zorder=3)
        ax2.plot(tab["time_ps"], tab["s_config"], color=color, lw=1.0,
                 ls="--", alpha=0.9, zorder=2)
    ax1.set_xscale("log")
    ax1.set_xlabel("Time (ps)", labelpad=3)
    ax1.set_ylabel(METRICS["s_atom"][1], labelpad=3)
    ax2.set_ylabel(METRICS["s_config"][1], labelpad=8, rotation=270)
    style_axes(ax1, hide_right=False)
    style_axes(ax2, hide_right=False)

    temp_handles = [Line2D([0], [0], color=temp_color(T), lw=1.4, label=f"{T} K")
                    for T in TEMPERATURES]
    style_handles = [
        Line2D([0], [0], color="0.25", lw=1.2, ls="-",
               label="$S_\\mathrm{atom}$ (left)"),
        Line2D([0], [0], color="0.25", lw=1.2, ls="--",
               label="$S_\\mathrm{config}$ (right)"),
    ]
    leg = ax1.legend(handles=temp_handles, frameon=False, loc="lower left",
                     handlelength=1.4, borderpad=0.3, labelspacing=0.25,
                     ncol=2, columnspacing=1.0)
    ax1.add_artist(leg)
    ax1.legend(handles=style_handles, frameon=False, loc="lower right",
               bbox_to_anchor=(1.0, 1.0), ncol=2, handlelength=1.8,
               borderpad=0.3, labelspacing=0.25, columnspacing=1.2)
    ax1.set_title(f"{system}  ({_WINDOW})", loc="left", pad=8)
    save(fig, f"{system}_entropy_dual")


def plot_overview(results):
    apply_style()
    fig, axes = plt.subplots(2, 2, figsize=(7.0, 4.6), sharex=True)
    for row, system in enumerate(SYSTEMS):
        for col, metric in enumerate(("s_atom", "s_config")):
            ax = axes[row, col]
            for temp_k in TEMPERATURES:
                tab = series(results, system, temp_k)
                if tab is None:
                    continue
                ax.plot(tab["time_ps"], tab[metric], color=temp_color(temp_k),
                        lw=1.0, label=f"{temp_k} K")
            ax.set_xscale("log")
            style_axes(ax)
            if row == 1:
                ax.set_xlabel("Time (ps)", labelpad=3)
            ax.set_ylabel(METRICS[metric][1], labelpad=3)
            ax.set_title(f"{system}  ({_WINDOW})", loc="left", pad=5)
            if row == 0 and col == 0:
                ax.legend(frameon=False, loc="best", ncol=2, handlelength=1.3,
                          labelspacing=0.2, columnspacing=0.9)
    fig.tight_layout(pad=0.6, w_pad=1.6)
    save(fig, "entropy_overview")


def write_long_csv(results):
    path = os.path.join(_OUT_DIR, "entropy_raw_long.csv")
    header = ("system,temperature_K,frame_index,time_ps,S_atom,S_config,"
              "n_motifs,n_clusters,n_metal,max_cluster_size")
    lines = [header]
    for system in SYSTEMS:
        for temp_k in TEMPERATURES:
            tab = series(results, system, temp_k)
            if tab is None:
                continue
            for r in tab:
                lines.append(
                    f"{system},{temp_k},{r['frame']},{r['time_ps']:.4f},"
                    f"{r['s_atom']:.6f},{r['s_config']:.6f},{r['n_motifs']},"
                    f"{r['n_clusters']},{r['n_metal']},{r['max_cluster']}"
                )
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    print(f"Saved: {path}")
    return path


def wide_columns(results, system, metric):
    """Per-temperature (time, value) column pairs; runs have unequal lengths."""
    columns = []
    for temp_k in TEMPERATURES:
        tab = series(results, system, temp_k)
        if tab is None:
            continue
        columns.append((f"time_ps_{temp_k}K", tab["time_ps"]))
        columns.append((f"{METRICS[metric][0]}_{temp_k}K", tab[metric]))
    return columns


def write_wide_csv(results, system, metric):
    columns = wide_columns(results, system, metric)
    if not columns:
        return None
    nrow = max(len(v) for _, v in columns)
    path = os.path.join(_OUT_DIR, f"{system}_{METRICS[metric][0]}_wide.csv")
    lines = [",".join(name for name, _ in columns)]
    for i in range(nrow):
        lines.append(",".join(
            f"{v[i]:.6f}" if i < len(v) else "" for _, v in columns
        ))
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    print(f"Saved: {path}")
    return path


def summary_rows(results):
    rows = [("system", "temperature_K", "n_atoms", "n_metal_atoms",
             "frames_in_dump", "frames_analysed", "t_max_ps",
             "S_atom_first", "S_atom_last10pct_mean", "S_config_first",
             "S_config_last10pct_mean", "n_clusters_first",
             "n_clusters_last", "max_cluster_first", "max_cluster_last")]
    for system in SYSTEMS:
        for temp_k in TEMPERATURES:
            payload = results.get((system, temp_k))
            if payload is None:
                continue
            tab = series(results, system, temp_k)
            tail = tab[int(0.9 * len(tab)):]
            rows.append((
                system, temp_k, int(payload["natoms"]), int(tab["n_metal"][0]),
                int(payload["n_frames_total"]), len(tab),
                round(float(tab["time_ps"].max()), 2),
                round(float(tab["s_atom"][0]), 4),
                round(float(tail["s_atom"].mean()), 4),
                round(float(tab["s_config"][0]), 4),
                round(float(tail["s_config"].mean()), 4),
                int(tab["n_clusters"][0]), int(tab["n_clusters"][-1]),
                int(tab["max_cluster"][0]), int(tab["max_cluster"][-1]),
            ))
    return rows


def write_summary_csv(results):
    rows = summary_rows(results)
    path = os.path.join(_OUT_DIR, "entropy_summary.csv")
    with open(path, "w", encoding="utf-8") as fh:
        for row in rows:
            fh.write(",".join(str(x) for x in row) + "\n")
    print(f"Saved: {path}")
    return path


def write_xlsx(results):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    readme = [
        ("Quantity", "Definition", "Units"),
        ("S_atom",
         "Shannon entropy over discrete local-environment motifs "
         f"(species, CN_O within {O_CUT} A, CN_metal within {MET_CUT} A); "
         "p_k = N_k / N_atoms", "nats"),
        ("S_config",
         "Shannon entropy over metal-cluster sizes; metal = all non-oxygen "
         f"atoms (Ni, Fe, Zn, Cr, Ru), clusters = connected components with "
         f"contact distance < {MET_CUT} A under PBC; p_c = n_c / N_metal",
         "nats"),
        ("n_motifs", "Number of distinct occupied motifs in the frame", "count"),
        ("n_clusters", "Number of metal clusters in the frame", "count"),
        ("max_cluster_size", "Atom count of the largest metal cluster", "count"),
        ("time_ps", "Simulation time from the dump Time attribute", "ps"),
        ("", "", ""),
        ("Sheet", "Contents", ""),
        ("summary", "One row per trajectory: sizes, lengths, start/end values", ""),
        ("raw_long", "Tidy raw data, one row per analysed frame", ""),
    ]
    for system in SYSTEMS:
        readme.append((f"{system[:20]}_Satom",
                       f"{system}: wide table, time/S_atom pairs per temperature", ""))
        readme.append((f"{system[:20]}_Sconf",
                       f"{system}: wide table, time/S_config pairs per temperature", ""))
    readme += [
        ("", "", ""),
        ("Method", "GPUMD NEP89 NVT-Langevin, dt = 1 fs, dump every 400 steps "
                   "(0.4 ps); 400 uniformly spaced frames analysed per run", ""),
        ("Window", f"Analysis window: {_WINDOW}", ""),
    ]
    for row in readme:
        ws.append(list(row))

    ws = wb.create_sheet("summary")
    for row in summary_rows(results):
        ws.append(list(row))

    ws = wb.create_sheet("raw_long")
    ws.append(["system", "temperature_K", "frame_index", "time_ps", "S_atom",
               "S_config", "n_motifs", "n_clusters", "n_metal",
               "max_cluster_size"])
    for system in SYSTEMS:
        for temp_k in TEMPERATURES:
            tab = series(results, system, temp_k)
            if tab is None:
                continue
            for r in tab:
                ws.append([system, temp_k, int(r["frame"]),
                           float(r["time_ps"]), float(r["s_atom"]),
                           float(r["s_config"]), int(r["n_motifs"]),
                           int(r["n_clusters"]), int(r["n_metal"]),
                           int(r["max_cluster"])])

    for system in SYSTEMS:
        for metric, tag in (("s_atom", "Satom"), ("s_config", "Sconf")):
            columns = wide_columns(results, system, metric)
            if not columns:
                continue
            ws = wb.create_sheet(f"{system[:20]}_{tag}"[:31])
            ws.append([name for name, _ in columns])
            nrow = max(len(v) for _, v in columns)
            for i in range(nrow):
                ws.append([float(v[i]) if i < len(v) else None
                           for _, v in columns])

    path = os.path.join(_OUT_DIR, "entropy_tables.xlsx")
    wb.save(path)
    print(f"Saved: {path}")
    return path


def build(results, out_dir: str | None = None, window_label: str | None = None):
    global _OUT_DIR, _WINDOW
    if out_dir is not None:
        _OUT_DIR = out_dir
    if window_label is not None:
        _WINDOW = window_label
    os.makedirs(_OUT_DIR, exist_ok=True)
    if not results:
        raise SystemExit("No results found; run the analysis first.")
    write_long_csv(results)
    write_summary_csv(results)
    for system in SYSTEMS:
        write_wide_csv(results, system, "s_atom")
        write_wide_csv(results, system, "s_config")
        plot_single(results, system, "s_atom")
        plot_single(results, system, "s_config")
        plot_dual(results, system)
    plot_overview(results)
    write_xlsx(results)


if __name__ == "__main__":
    from run_entropy import load_results

    build(load_results())
